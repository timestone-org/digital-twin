"""Word 版式、富文本、降级与位图后备的真实文件测试。"""

import io
import zipfile
from datetime import UTC, datetime

import pytest
from docx import Document

from platform_server.apps.report.schemas.document import (
    DocumentNode,
    PageSettings,
    Watermark,
)
from platform_server.apps.report.schemas.preview import (
    ChartSeries,
    NodeValue,
    PreviewOut,
    SeriesPoint,
)
from platform_server.apps.report.schemas.template import TemplateBody
from platform_server.apps.report.services.chart_raster import render_chart
from platform_server.apps.report.services.docx_builder import build_docx
from platform_server.apps.report.services.docx_import import validate_archive
from platform_server.apps.report.services.word_task import WordTask, run_task


def empty_preview(nodes=None):
    return PreviewOut(
        is_valid=True,
        period="2026-08",
        timezone="UTC",
        metrics=[],
        warnings=[],
        nodes=nodes or {},
    )


def test_word_page_orientation_watermark_and_toc():
    page = PageSettings(
        orientation="landscape",
        header="保留页眉",
        footer="保留页脚",
        watermark=Watermark(text="内部资料"),
        is_toc_enabled=True,
    )
    result = build_docx(
        TemplateBody(
            name="报告",
            page_json=page,
            doc_json=DocumentNode(
                type="doc",
                content=[
                    DocumentNode(
                        type="heading",
                        attrs={"level": 2},
                        content=[DocumentNode(type="text", text="内容")],
                    )
                ],
            ),
        ),
        empty_preview(),
    )
    document = Document(io.BytesIO(result.payload))
    assert document.sections[0].page_width > document.sections[0].page_height
    assert "保留页眉" in document.sections[0].header.paragraphs[0].text
    with zipfile.ZipFile(io.BytesIO(result.payload)) as archive:
        assert b"ReportWatermark" in archive.read("word/header1.xml")
        assert b"_x0000_t136" in archive.read("word/header1.xml")
        assert b"TOC" in archive.read("word/document.xml")
        assert b"updateFields" in archive.read("word/settings.xml")


def test_invalid_page_margins_fail_explicitly():
    page = PageSettings(size="custom", width_cm=5, height_cm=5)
    with pytest.raises(ValueError, match="页边距"):
        build_docx(TemplateBody(name="报告", page_json=page), empty_preview())


def test_rich_text_marks_and_unsupported_node_warnings():
    node = DocumentNode.model_validate(
        {
            "type": "doc",
            "content": [
                {
                    "type": "paragraph",
                    "content": [
                        {
                            "type": "text",
                            "text": "样式",
                            "marks": [
                                {"type": "bold"},
                                {"type": "italic"},
                                {"type": "underline"},
                                {"type": "strike"},
                                {
                                    "type": "textStyle",
                                    "attrs": {
                                        "fontFamily": "Arial",
                                        "fontSize": "16pt",
                                        "color": "#112233",
                                    },
                                },
                            ],
                        },
                        {"type": "hardBreak"},
                        {"type": "metricRef"},
                        {"type": "unsupported"},
                    ],
                },
                {"type": "pageBreak"},
                {"type": "unknown", "text": "保留文字"},
            ],
        }
    )
    result = build_docx(
        TemplateBody(name="报告", doc_json=node), empty_preview()
    )
    document = Document(io.BytesIO(result.payload))
    run = document.paragraphs[0].runs[0]
    assert run.bold
    assert run.italic
    assert run.underline
    assert run.font.strike
    assert run.font.size.pt == 16
    assert str(run.font.color.rgb) == "112233"
    assert len(result.warnings) == 2
    assert "保留文字" in [paragraph.text for paragraph in document.paragraphs]


@pytest.mark.parametrize("kind", ["line", "bar"])
def test_raster_charts_are_real_png(kind):
    node = NodeValue(
        kind=kind,
        title="能耗",
        series=[
            ChartSeries(
                name="能耗",
                points=[
                    SeriesPoint(
                        ts=datetime(2026, 8, day, tzinfo=UTC), value=value
                    )
                    for day, value in [(1, "10"), (2, None), (3, "12")]
                ],
            )
        ],
    )
    result = render_chart(node, "Asia/Shanghai")
    assert result.startswith(b"\x89PNG\r\n\x1a\n")
    assert len(result) > 1000


def test_explicit_raster_and_empty_business_blocks():
    node = DocumentNode(
        type="doc",
        content=[
            DocumentNode(type="dsChart", attrs={"render": "raster"}),
            DocumentNode(type="dsTable"),
            DocumentNode(type="dsChart"),
        ],
    )
    chart = NodeValue(
        kind="line",
        title="趋势",
        series=[
            ChartSeries(
                name="值",
                points=[
                    SeriesPoint(ts=datetime(2026, 8, 1, tzinfo=UTC), value="1")
                ],
            )
        ],
    )
    preview = empty_preview(
        {"content.0": chart, "content.1": NodeValue(kind="dsTable")}
    )
    result = build_docx(TemplateBody(name="报告", doc_json=node), preview)
    document = Document(io.BytesIO(result.payload))
    assert len(document.inline_shapes) == 1
    assert "暂无数据" in [paragraph.text for paragraph in document.paragraphs]
    assert "[数据未配置]" in [
        paragraph.text for paragraph in document.paragraphs
    ]


def test_word_task_import_and_invalid_input():
    document = Document()
    document.add_heading("导入标题", level=1)
    source = io.BytesIO()
    document.save(source)
    result = run_task(WordTask(source=source.getvalue()))
    assert result.imported is not None
    assert result.imported.doc_json.content[0].type == "heading"
    with pytest.raises(ValueError, match="缺少"):
        run_task(WordTask())


def test_upload_archive_limits():
    with pytest.raises(ValueError, match="超过"):
        validate_archive(b"x" * (10 * 1024 * 1024 + 1))
    source = io.BytesIO()
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("other.xml", "")
    with pytest.raises(ValueError, match="有效"):
        validate_archive(source.getvalue())
