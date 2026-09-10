"""报告字体、字号单位、业务行内样式与语义排版的导出契约。"""

import io
from datetime import UTC, datetime, timedelta

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from platform_server.apps.report.schemas.document import (
    DocumentMark,
    DocumentNode,
    PageSettings,
)
from platform_server.apps.report.schemas.preview import (
    ChartSeries,
    NodeValue,
    PreviewOut,
    SeriesPoint,
)
from platform_server.apps.report.schemas.template import TemplateBody
from platform_server.apps.report.services import docx_builder
from platform_server.apps.report.services.docx_builder import build_docx


def render(nodes, *, page=None, values=None):
    template = TemplateBody(
        name="运行月报",
        page_json=page or PageSettings(),
        doc_json=DocumentNode(type="doc", content=nodes),
    )
    preview = PreviewOut(
        is_valid=True,
        period="2026-08",
        timezone="UTC",
        metrics=[],
        nodes=values or {},
        warnings=[],
    )
    return Document(io.BytesIO(build_docx(template, preview).payload))


def test_body_and_headings_set_chinese_fonts_without_theme_overrides():
    document = render([], page=PageSettings(font_family="Noto Sans CJK SC"))
    for name in (
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "Caption",
        "Header",
        "Footer",
    ):
        style = document.styles[name]
        fonts = style.element.rPr.rFonts
        assert fonts.get(qn("w:eastAsia")) == "Noto Sans CJK SC"
        assert fonts.get(qn("w:eastAsiaTheme")) is None
        assert fonts.get(qn("w:asciiTheme")) is None
    assert document.styles["Normal"].paragraph_format.line_spacing == Pt(19.8)
    assert str(document.styles["Heading 1"].font.color.rgb) == "000000"


def test_pixels_convert_to_points_and_business_runs_keep_marks():
    marks = [
        DocumentMark(
            type="textStyle",
            attrs={"fontSize": "16px", "fontFamily": "Noto Sans CJK SC"},
        ),
        DocumentMark(type="bold"),
    ]
    paragraph = DocumentNode(
        type="paragraph",
        content=[
            DocumentNode(type="text", text="温度：", marks=marks),
            DocumentNode(type="metricRef", marks=marks),
            DocumentNode(type="condText", marks=marks),
        ],
    )
    document = render(
        [paragraph],
        values={
            "content.0.content.1": NodeValue(kind="metricRef", text="22.52"),
            "content.0.content.2": NodeValue(kind="condText", text="运行正常"),
        },
    )
    for run in document.paragraphs[0].runs:
        assert run.font.size.pt == 12
        assert run.bold
        assert (
            run.element.rPr.rFonts.get(qn("w:eastAsia")) == "Noto Sans CJK SC"
        )


def test_title_precedes_a_populated_toc_and_headings_stay_with_content():
    nodes = [
        DocumentNode(
            type="heading",
            attrs={"level": level},
            content=[DocumentNode(type="text", text=text)],
        )
        for level, text in [(1, "运行月报"), (2, "核心指标")]
    ]
    document = render(nodes, page=PageSettings(is_toc_enabled=True))
    assert document.paragraphs[0].style.name == "Title"
    assert document.paragraphs[0].text == "运行月报"
    assert "核心指标" in [p.text for p in document.paragraphs[1:-1]]
    assert document.styles["Heading 2"].paragraph_format.keep_with_next
    assert document.styles["Title"].element.pPr.find(qn("w:pBdr")) is None


def test_tables_repeat_header_keep_rows_and_align_numeric_values():
    document = render(
        [DocumentNode(type="dsTable")],
        values={
            "content.0": NodeValue(
                kind="dsTable",
                columns=["时间", "温度"],
                rows=[["2026-08-01", "22.52"], ["2026-08-02", "23.00"]],
            )
        },
    )
    table = document.tables[0]
    assert (
        table.rows[0]._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None
    )
    assert len(table.rows[0]._tr.findall(qn("w:trPr"))) == 1
    assert (
        table.rows[1]._tr.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None
    )
    assert table.cell(0, 0).paragraphs[0].runs[0].bold
    assert table.cell(1, 1).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert table.cell(1, 1).paragraphs[0].runs[0].font.size == Pt(10.5)
    assert table.cell(1, 1).text == "22.52"


@pytest.mark.parametrize("raster_available", [True, False])
def test_dense_chart_uses_readable_image_with_native_fallback(
    monkeypatch, raster_available
):
    if not raster_available:

        def unavailable(*_args):
            raise ValueError("字体不可用")

        monkeypatch.setattr(docx_builder, "render_chart", unavailable)
    points = [
        SeriesPoint(
            ts=datetime(2026, 8, 1, tzinfo=UTC) + timedelta(hours=index),
            value=str(index),
        )
        for index in range(13)
    ]
    document = render(
        [DocumentNode(type="dsChart")],
        values={
            "content.0": NodeValue(
                kind="line",
                title="温度趋势",
                series=[ChartSeries(name="温度", points=points)],
            )
        },
    )
    xml = document.element.xml
    assert ("<pic:pic" in xml) == raster_available
    assert ("<c:chart" in xml) != raster_available
    assert len(document.inline_shapes) == 1
