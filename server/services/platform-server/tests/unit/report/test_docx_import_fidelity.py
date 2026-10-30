"""Word 字体、段落、表格及链接显示文本的导入出稿契约。"""

import io
import zipfile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from unit.report.docx_import_fakes import (
    fixture_document,
    payload,
    render_import,
    roundtrip,
)

from platform_server.apps.report.schemas.document import DocumentNode
from platform_server.apps.report.schemas.template import TemplateBody
from platform_server.apps.report.services.document import validate_template
from platform_server.apps.report.services.docx_import import import_docx


def test_import_resolves_style_font_size_and_basic_run_marks():
    imported, generated, _result = roundtrip(fixture_document())
    paragraph = imported.doc_json.content[0]
    mark = next(
        mark for mark in paragraph.content[0].marks if mark.type == "textStyle"
    )
    assert mark.attrs["fontFamily"] == "Noto Serif CJK SC"
    assert mark.attrs["fontSize"] == "13.5pt"
    assert "bold" in {mark.type for mark in paragraph.content[0].marks}
    assert generated.paragraphs[0].runs[0].font.size == Pt(13.5)
    second = generated.paragraphs[0].runs[1]
    assert second.font.size == Pt(10.5)
    assert second.italic
    assert second.underline
    assert second.font.strike
    assert str(second.font.color.rgb) == "445566"
    assert (
        second.font.element.rPr.rFonts.get(qn("w:eastAsia"))
        == "Noto Sans CJK SC"
    )


def test_import_roundtrip_preserves_effective_paragraph_layout():
    imported, generated, _result = roundtrip(fixture_document())
    attrs = imported.doc_json.content[0].attrs
    assert attrs["textAlign"] == "right"
    assert attrs["firstLineIndentPt"] == 21
    assert attrs["margin"] == {"top": 8, "bottom": 12}
    paragraph = generated.paragraphs[0]
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert paragraph.paragraph_format.left_indent == Pt(12)
    assert paragraph.paragraph_format.right_indent == Pt(8)
    assert paragraph.paragraph_format.first_line_indent == Pt(21)
    assert paragraph.paragraph_format.space_before == Pt(6)
    assert paragraph.paragraph_format.space_after == Pt(9)
    assert paragraph.paragraph_format.line_spacing == 1.5


@pytest.mark.parametrize(
    "color", ["rgb(17, 34, 51)", "rgba(17, 34, 51, 1)", "#123"]
)
def test_browser_normalized_text_color_remains_in_word(color):
    imported = import_docx(payload(fixture_document()))
    first = imported.doc_json.content[0].content[0]
    style = next(mark for mark in first.marks if mark.type == "textStyle")
    style.attrs["color"] = color
    generated, _result = render_import(imported)
    assert str(generated.paragraphs[0].runs[0].font.color.rgb) == "112233"


def test_explicit_disabled_heading_emphasis_has_a_fidelity_warning():
    document = Document()
    paragraph = document.add_paragraph(style="Heading 2")
    paragraph.add_run("取消标题加粗").bold = False
    imported = import_docx(payload(document))
    assert any("继承" in warning for warning in imported.dropped)


def test_inherited_title_border_and_direct_shading_have_precise_loss_warnings():
    document = Document()
    paragraph = document.add_paragraph("继承标题边框", style="Title")
    properties = document.styles["Title"].element.get_or_add_pPr()
    if properties.find(qn("w:pBdr")) is None:
        properties.append(OxmlElement("w:pBdr"))
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:shd"))
    imported = import_docx(payload(document))
    assert "段落边框未导入" in imported.dropped
    assert "段落底纹未导入" in imported.dropped


def test_document_default_font_size_and_spacing_are_not_replaced():
    document = Document()
    document.add_paragraph("文档默认样式")
    imported, generated, _result = roundtrip(document)
    assert imported.doc_json.content[0].content[0].marks
    assert generated.paragraphs[0].runs[0].font.size == Pt(11)
    assert generated.paragraphs[0].paragraph_format.space_after == Pt(10)
    assert generated.paragraphs[0].paragraph_format.line_spacing == 1.15


def test_real_editor_string_spacing_and_native_indent_remain_compatible():
    imported = import_docx(payload(fixture_document()))
    attrs = imported.doc_json.content[0].attrs
    attrs.pop("firstLineIndentPt")
    attrs.update(
        {
            "indent": 10.5,
            "indentUnit": "pt",
            "lineHeight": "1.5",
            "margin": {"top": "8", "bottom": "12"},
        }
    )
    document, _result = render_import(imported)
    formatting = document.paragraphs[0].paragraph_format
    assert formatting.first_line_indent == Pt(21)
    assert formatting.space_before == Pt(6)
    assert formatting.space_after == Pt(9)
    assert formatting.line_spacing == 1.5


@pytest.mark.parametrize("header", [False, True])
def test_only_explicit_table_header_rows_repeat_across_pages(header):
    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格第一行"
    if header:
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    imported, generated, _result = roundtrip(document)
    assert ("w:tblHeader" in generated.tables[0]._tbl.xml) is header
    assert any("列宽" in warning for warning in imported.dropped)


def test_expanded_static_table_grid_has_an_upfront_memory_budget():
    table = DocumentNode(
        type="table",
        content=[
            DocumentNode(
                type="tableRow",
                content=[
                    DocumentNode(type="tableCell", attrs={"colspan": 100})
                ],
            )
            for _index in range(1000)
        ],
    )
    checked = validate_template(
        TemplateBody(
            name="展开预算", doc_json=DocumentNode(type="doc", content=[table])
        )
    )
    assert not checked.is_valid
    assert any("10000" in issue.message for issue in checked.issues)


def test_static_table_spans_cannot_overlap_an_existing_vertical_merge():
    table = DocumentNode(
        type="table",
        content=[
            DocumentNode(
                type="tableRow",
                content=[
                    DocumentNode(type="tableCell"),
                    DocumentNode(type="tableCell", attrs={"rowspan": 2}),
                ],
            ),
            DocumentNode(
                type="tableRow",
                content=[DocumentNode(type="tableCell", attrs={"colspan": 2})],
            ),
        ],
    )
    checked = validate_template(
        TemplateBody(
            name="跨度冲突", doc_json=DocumentNode(type="doc", content=[table])
        )
    )
    assert not checked.is_valid
    assert any("重叠" in issue.message for issue in checked.issues)


@pytest.mark.parametrize("indent", [-21, 0, 21])
def test_first_line_indentation_has_no_two_point_rounding(indent):
    document = Document()
    paragraph = document.add_paragraph("首行与悬挂缩进")
    paragraph.paragraph_format.first_line_indent = Pt(indent)
    _imported, result, _generated = roundtrip(document)
    assert result.paragraphs[0].paragraph_format.first_line_indent == Pt(indent)


def test_exact_line_spacing_and_zero_margins_survive_roundtrip():
    document = Document()
    paragraph = document.add_paragraph("固定行距")
    paragraph.paragraph_format.line_spacing = Pt(22)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    _imported, result, _generated = roundtrip(document)
    target = result.paragraphs[0].paragraph_format
    assert target.line_spacing_rule == WD_LINE_SPACING.EXACTLY
    assert target.line_spacing == Pt(22)
    assert target.space_before == Pt(0)
    assert target.space_after == Pt(0)


def test_table_cell_layout_does_not_get_overwritten_by_report_defaults():
    document = Document()
    table = document.add_table(rows=1, cols=1)
    paragraph = table.cell(0, 0).paragraphs[0]
    paragraph.add_run("单元格格式").font.size = Pt(14)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = Pt(24)
    _imported, result, _generated = roundtrip(document)
    target = result.tables[0].cell(0, 0).paragraphs[0]
    assert target.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert target.paragraph_format.space_before == Pt(3)
    assert target.paragraph_format.space_after == Pt(7)
    assert target.paragraph_format.line_spacing == Pt(24)
    assert target.runs[0].font.size == Pt(14)


def test_merged_table_has_single_content_and_native_span_structure():
    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(1, 1)).text = "合并内容"
    table.cell(0, 2).text = "右上"
    table.cell(1, 2).text = "右下"
    imported, result, _generated = roundtrip(document)
    cells = imported.doc_json.content[0].content[0].content
    assert cells[0].attrs["colspan"] == 2
    assert cells[0].attrs["rowspan"] == 2
    assert len(cells) == 2
    assert len(imported.doc_json.content[0].content[1].content) == 1
    assert result.tables[0].cell(0, 0)._tc is result.tables[0].cell(1, 1)._tc
    xml = result.tables[0]._tbl.xml
    assert xml.count("合并内容") == 1


def test_breaks_and_external_hyperlink_display_text_are_not_lost():
    document = Document()
    paragraph = document.add_paragraph("之前")
    relation = document.part.relate_to(
        "https://example.test/private",
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "显示链接文字"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    paragraph.add_run().add_break(WD_BREAK.LINE)
    paragraph.add_run("之后")
    imported, result, generated = roundtrip(document)
    assert result.paragraphs[0].text == "之前显示链接文字\n之后"
    assert any("超链接" in warning for warning in imported.dropped)
    with zipfile.ZipFile(io.BytesIO(generated.payload)) as archive:
        assert b"https://example.test" not in archive.read(
            "word/_rels/document.xml.rels"
        )


def test_unmappable_formatting_is_reported_instead_of_silent_loss():
    document = Document()
    paragraph = document.add_paragraph("最小行距与制表位")
    paragraph.paragraph_format.line_spacing = Pt(20)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(100))
    imported = import_docx(payload(document))
    assert any("最小行距" in warning for warning in imported.dropped)
    assert any("制表位" in warning for warning in imported.dropped)


def test_macro_payload_is_rejected_before_parser_runs():
    base = payload(Document())
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(base)) as source,
        zipfile.ZipFile(output, "w") as target,
    ):
        for entry in source.infolist():
            target.writestr(entry, source.read(entry))
        target.writestr("word/vbaProject.bin", b"inactive-test-fixture")
    with pytest.raises(ValueError, match="宏"):
        import_docx(output.getvalue())


def test_theme_font_projection_reports_static_font_conversion():
    document = Document()
    document.add_paragraph("主题字体的中西文字形需核对", style="Title")
    imported = import_docx(payload(document))
    assert any("主题字体" in warning for warning in imported.dropped)
