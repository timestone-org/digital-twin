"""Word 导入测试的真实文档与字节往返假件。"""

import io

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from platform_server.apps.report.schemas.preview import PreviewOut
from platform_server.apps.report.schemas.template import TemplateBody
from platform_server.apps.report.services.docx_builder import build_docx
from platform_server.apps.report.services.docx_import import import_docx


def payload(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def fixture_document():
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.25
    style = document.styles.add_style("报告正文", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = normal
    style.font.name = "Noto Serif CJK SC"
    style.font.size = Pt(13.5)
    style.font.bold = True
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style.paragraph_format.left_indent = Pt(12)
    style.paragraph_format.right_indent = Pt(8)
    style.paragraph_format.first_line_indent = Pt(21)
    style.paragraph_format.space_before = Pt(6)
    style.paragraph_format.space_after = Pt(9)
    style.paragraph_format.line_spacing = 1.5
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run("继承样式中文")
    run = paragraph.add_run("文字样式")
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string("445566")
    run.font.italic = True
    run.font.underline = True
    run.font.strike = True
    run.font.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    return document


def roundtrip(document):
    imported = import_docx(payload(document))
    generated_document, generated = render_import(imported)
    return imported, generated_document, generated


def render_import(imported):
    preview = PreviewOut(
        is_valid=True,
        period="2026-10",
        timezone="UTC",
        metrics=[],
        nodes={},
        warnings=[],
    )
    generated = build_docx(
        TemplateBody(
            name="Word 导入核对",
            doc_json=imported.doc_json,
            page_json=imported.page_json,
        ),
        preview,
    )
    return Document(io.BytesIO(generated.payload)), generated
