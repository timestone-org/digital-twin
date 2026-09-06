"""实际 Word 文件的文本、表格、图表与导入往返。"""

import io
import zipfile
from datetime import UTC, datetime

from docx import Document
from docx.oxml import parse_xml

from platform_server.apps.report.schemas.document import DocumentNode
from platform_server.apps.report.schemas.preview import (
    ChartSeries,
    NodeValue,
    PreviewOut,
    SeriesPoint,
)
from platform_server.apps.report.schemas.template import TemplateBody
from platform_server.apps.report.services.docx_builder import build_docx
from platform_server.apps.report.services.docx_import import import_docx


def test_word_contains_resolved_text_and_native_chart():
    template = chart_template()
    preview = PreviewOut(
        is_valid=True,
        period="2026-08",
        timezone="Asia/Shanghai",
        metrics=[],
        nodes={
            "content.1.content.1": NodeValue(
                kind="metricRef", text="123.45 kWh"
            ),
            "content.2": NodeValue(
                kind="bar",
                title="能耗",
                series=[
                    ChartSeries(
                        name="能耗",
                        points=[
                            SeriesPoint(
                                ts=datetime(2026, 8, 1, tzinfo=UTC),
                                value="123.45",
                            )
                        ],
                    )
                ],
            ),
        },
        warnings=[],
    )
    result = build_docx(template, preview)
    document = Document(io.BytesIO(result.payload))
    assert "本期能耗：123.45 kWh" in [
        paragraph.text for paragraph in document.paragraphs
    ]
    with zipfile.ZipFile(io.BytesIO(result.payload)) as archive:
        chart = parse_xml(archive.read("word/charts/chart1.xml"))
        assert (
            chart.find(
                ".//{http://schemas.openxmlformats.org/drawingml/2006/chart}barChart"
            )
            is not None
        )
    assert not result.warnings


def test_word_import_preserves_paragraph_order_with_table():
    document = Document()
    document.add_paragraph("之前")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "单元格"
    document.add_paragraph("之后")
    payload = io.BytesIO()
    document.save(payload)
    result = import_docx(payload.getvalue())
    assert [node.type for node in result.doc_json.content] == [
        "paragraph",
        "table",
        "paragraph",
    ]
    assert (
        result.doc_json.content[1]
        .content[0]
        .content[0]
        .content[0]
        .content[0]
        .text
        == "单元格"
    )


def test_table_inline_metric_has_no_extra_empty_paragraph():
    template = inline_table_template()
    preview = PreviewOut(
        is_valid=True,
        period="2026-08",
        timezone="UTC",
        metrics=[],
        warnings=[],
        nodes={
            "content.0.content.0.content.0.content.0.content.1": NodeValue(
                kind="metricRef", text="12.30"
            )
        },
    )
    result = build_docx(template, preview)
    document = Document(io.BytesIO(result.payload))
    assert document.tables[0].cell(0, 0).text == "能耗 12.30"
    with zipfile.ZipFile(io.BytesIO(result.payload)) as archive:
        assert b"tblHeader" in archive.read("word/document.xml")


def test_list_semantics_survive_word_export():
    template = TemplateBody(
        name="列表报告",
        doc_json=DocumentNode.model_validate(
            {
                "type": "doc",
                "content": [
                    {
                        "type": "bulletList",
                        "content": [
                            {
                                "type": "listItem",
                                "content": [
                                    {
                                        "type": "paragraph",
                                        "content": [
                                            {"type": "text", "text": "检查能耗"}
                                        ],
                                    }
                                ],
                            }
                        ],
                    }
                ],
            }
        ),
    )
    preview = PreviewOut(
        is_valid=True,
        period="2026-08",
        timezone="UTC",
        metrics=[],
        warnings=[],
        nodes={},
    )
    document = Document(io.BytesIO(build_docx(template, preview).payload))
    assert document.paragraphs[0].style.name == "List Bullet"


def chart_template():
    return TemplateBody(
        name="能耗月报",
        doc_json=DocumentNode(
            type="doc",
            content=[
                DocumentNode(
                    type="heading",
                    attrs={"level": 1},
                    content=[DocumentNode(type="text", text="能耗月报")],
                ),
                DocumentNode(
                    type="paragraph",
                    content=[
                        DocumentNode(type="text", text="本期能耗："),
                        DocumentNode(
                            type="metricRef", attrs={"expr": "{能耗}"}
                        ),
                    ],
                ),
                DocumentNode(type="dsChart"),
            ],
        ),
    )


def inline_table_template():
    paragraph = DocumentNode(
        type="paragraph",
        content=[
            DocumentNode(type="text", text="能耗 "),
            DocumentNode(type="metricRef"),
        ],
    )
    cell = DocumentNode(type="tableCell", content=[paragraph])
    row = DocumentNode(type="tableRow", content=[cell])
    table = DocumentNode(type="table", content=[row])
    return TemplateBody(
        name="表格报告", doc_json=DocumentNode(type="doc", content=[table])
    )
