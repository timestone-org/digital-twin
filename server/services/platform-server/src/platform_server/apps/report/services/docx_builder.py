"""报告 AST 到 Word 的唯一生成器，运行于 worker 子进程。"""

import io
from dataclasses import dataclass, field
from typing import Protocol

from docx import Document
from docx.document import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from platform_server.apps.report.schemas.document import (
    DocumentMark,
    DocumentNode,
)
from platform_server.apps.report.schemas.preview import NodeValue, PreviewOut
from platform_server.apps.report.schemas.template import TemplateBody
from platform_server.apps.report.services.chart_ooxml import add_chart
from platform_server.apps.report.services.chart_raster import render_chart
from platform_server.apps.report.services.document import text_attr
from platform_server.apps.report.services.docx_page import apply_page
from platform_server.apps.report.services.ooxml import append_xml

_HEX_COLOR_LENGTH = 7


class BlockParent(Protocol):
    """文档或单元格的块级写面。"""

    def add_paragraph(self, text: str = "") -> Paragraph: ...
    def add_table(self, rows: int, cols: int) -> Table: ...

    @property
    def paragraphs(self) -> list[Paragraph]: ...


@dataclass
class RenderState:
    """一份文档的生成上下文。"""

    document: WordDocument
    preview: PreviewOut
    warnings: list[str] = field(default_factory=list[str])
    list_style: str | None = None


@dataclass(frozen=True)
class WordResult:
    """成品与降级清单。"""

    payload: bytes
    warnings: tuple[str, ...]


def build_docx(template: TemplateBody, preview: PreviewOut) -> WordResult:
    """从模板和冻结的试算结果生成 Word。Args: template, preview。"""
    document = Document()
    apply_page(document, template.page_json)
    state = RenderState(
        document=document, preview=preview, warnings=list(preview.warnings)
    )
    for index, node in enumerate(template.doc_json.content):
        _block(state, document, node, f"content.{index}")
    if state.warnings:
        document.add_heading("数据与版式提示", level=2)
        for warning in dict.fromkeys(state.warnings):
            document.add_paragraph(warning)
    output = io.BytesIO()
    document.save(output)
    return WordResult(
        payload=output.getvalue(), warnings=tuple(dict.fromkeys(state.warnings))
    )


def _block(
    state: RenderState, parent: BlockParent, node: DocumentNode, path: str
) -> None:
    if node.type in ("paragraph", "heading"):
        paragraph = parent.add_paragraph()
        if state.list_style:
            paragraph.style = state.list_style
        _set_paragraph_style(paragraph, node)
        _paragraph(state, paragraph, node, path)
        return
    if node.type == "table":
        _table(state, parent, node, path)
        return
    if node.type == "pageBreak":
        parent.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        return
    if node.type in ("dsTable", "dsChart"):
        _business_block(state, parent, node, path)
        return
    if node.type in (
        "bulletList",
        "orderedList",
        "listItem",
        "blockquote",
        "doc",
    ):
        _container_block(state, parent, node, path)
        return
    state.warnings.append(f"{path}：不支持的节点 {node.type}")
    if node.text:
        parent.add_paragraph(node.text)


def _container_block(
    state: RenderState, parent: BlockParent, node: DocumentNode, path: str
) -> None:
    previous = state.list_style
    if node.type in ("bulletList", "orderedList"):
        state.list_style = (
            "List Bullet" if node.type == "bulletList" else "List Number"
        )
    try:
        for index, child in enumerate(node.content):
            _block(state, parent, child, f"{path}.content.{index}")
    finally:
        state.list_style = previous


def _paragraph(
    state: RenderState, paragraph: Paragraph, node: DocumentNode, path: str
) -> None:
    alignments = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    paragraph.alignment = alignments.get(
        text_attr(node, "textAlign", "left"), WD_ALIGN_PARAGRAPH.LEFT
    )
    for index, child in enumerate(node.content):
        child_path = f"{path}.content.{index}"
        if child.type == "text":
            _marks(paragraph.add_run(child.text or ""), child)
        elif child.type == "hardBreak":
            paragraph.add_run().add_break()
        elif child.type in ("metricRef", "condText"):
            value = state.preview.nodes.get(child_path)
            paragraph.add_run(value.text if value else "[数据未配置]")
        else:
            state.warnings.append(
                f"{child_path}：行内节点 {child.type} 无法保留"
            )


def _marks(run: Run, node: DocumentNode) -> None:
    for mark in node.marks:
        if mark.type in ("bold", "italic", "underline", "strike"):
            target = run.font if mark.type == "strike" else run
            setattr(target, mark.type, True)
        elif mark.type == "textStyle":
            _text_style(run, mark)


def _set_paragraph_style(paragraph: Paragraph, node: DocumentNode) -> None:
    if node.type == "heading":
        level = node.attrs.get("level", 1)
        heading_level = max(1, min(6, level)) if isinstance(level, int) else 1
        paragraph.style = f"Heading {heading_level}"


def _text_style(run: Run, mark: DocumentMark) -> None:
    font = mark.attrs.get("fontFamily")
    if isinstance(font, str):
        run.font.name = font
    size = mark.attrs.get("fontSize")
    if isinstance(size, str):
        number = size.removesuffix("pt").removesuffix("px")
        if number.replace(".", "", 1).isdigit():
            run.font.size = Pt(min(72, max(6, float(number))))
    color = mark.attrs.get("color")
    if (
        isinstance(color, str)
        and len(color) == _HEX_COLOR_LENGTH
        and color.startswith("#")
    ):
        _set_color(run, color)


def _set_color(run: Run, color: str) -> None:
    try:
        run.font.color.rgb = RGBColor.from_string(color[1:])
    except ValueError:
        return


def _table(
    state: RenderState,
    parent: BlockParent,
    node: DocumentNode,
    path: str,
) -> None:
    rows = node.content
    if not rows:
        return
    columns = max((len(row.content) for row in rows), default=1)
    table = parent.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index, cell in enumerate(row.content):
            target = table.cell(row_index, column_index)
            cell_path = f"{path}.content.{row_index}.content.{column_index}"
            _cell_blocks(state, target, cell, cell_path)
    _repeat_header(table.rows[0])


def _cell_blocks(
    state: RenderState, target: BlockParent, cell: DocumentNode, path: str
) -> None:
    for index, child in enumerate(cell.content):
        child_path = f"{path}.content.{index}"
        if index == 0 and child.type in ("paragraph", "heading"):
            paragraph = target.paragraphs[0]
            _set_paragraph_style(paragraph, child)
            _paragraph(state, paragraph, child, child_path)
        else:
            _block(state, target, child, child_path)


def _business_block(
    state: RenderState,
    parent: BlockParent,
    node: DocumentNode,
    path: str,
) -> None:
    value = state.preview.nodes.get(path)
    if value is None or value.text:
        parent.add_paragraph(value.text if value else "[数据未配置]")
        return
    if value.title:
        parent.add_paragraph(value.title)
    if node.type == "dsChart":
        if not isinstance(parent, WordDocument):
            state.warnings.append(f"{path}：表格内图表未生成")
            return
        _draw_chart(state, parent, node, value, path)
    else:
        _data_table(parent, value)


def _draw_chart(
    state: RenderState,
    document: WordDocument,
    node: DocumentNode,
    value: NodeValue,
    path: str,
) -> None:
    if text_attr(node, "render", "native") == "native":
        try:
            add_chart(document, value, state.preview.timezone)
            return
        except ValueError:
            state.warnings.append(f"{path}：原生图表失败，已尝试位图")
    try:
        payload = render_chart(value, state.preview.timezone)
        document.add_picture(io.BytesIO(payload), width=Cm(14))
    except (ValueError, RuntimeError):
        state.warnings.append(f"{path}：位图生成失败，请检查中文字体")
        document.add_paragraph("[图表不可用]")


def _data_table(parent: BlockParent, value: NodeValue) -> None:
    if not value.columns:
        parent.add_paragraph("暂无数据")
        return
    table = parent.add_table(rows=1, cols=len(value.columns))
    table.style = "Table Grid"
    for index, name in enumerate(value.columns):
        table.rows[0].cells[index].text = name
    _repeat_header(table.rows[0])
    for row in value.rows:
        cells = table.add_row().cells
        for index, text in enumerate(row[: len(cells)]):
            cells[index].text = text


def _repeat_header(row: object) -> None:
    append_xml(
        row,
        '<w:trPr xmlns:w="http://schemas.openxmlformats.org/'
        'wordprocessingml/2006/main"><w:tblHeader/></w:trPr>',
        attribute="_tr",
    )
