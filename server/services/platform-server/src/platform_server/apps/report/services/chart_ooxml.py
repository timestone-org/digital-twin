"""原生 DrawingML 折线和柱状图，复用备份模块的缓存数据与 OPC 装配方式。"""

from math import ceil
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from docx.document import Document
from docx.opc.part import Part
from docx.oxml import parse_xml

from platform_server.apps.report.schemas.preview import ChartSeries, NodeValue
from platform_server.apps.report.services.aggregation import decimal_value
from platform_server.apps.report.services.ooxml import append_xml

C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CHART_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
)
CHART_REL = f"{R_NS}/chart"
_CAT_AXIS = 101
_VAL_AXIS = 102


def _series_xml(
    series: ChartSeries,
    index: int,
    categories: list[str],
    stamps: list[str],
    kind: str,
) -> str:
    values = {
        point.ts.isoformat(): decimal_value(point.value)
        for point in series.points
    }
    names = "".join(
        f'<c:pt idx="{index}"><c:v>{escape(name)}</c:v></c:pt>'
        for index, name in enumerate(categories)
    )
    numbers = "".join(
        f'<c:pt idx="{index}"><c:v>{values[stamp]}</c:v></c:pt>'
        for index, stamp in enumerate(stamps)
        if values.get(stamp) is not None
    )
    return (
        f'<c:ser><c:idx val="{index}"/><c:order val="{index}"/>'
        f"<c:tx><c:v>{escape(series.name)}</c:v></c:tx>"
        f"{_series_style(index, kind)}"
        f"<c:cat><c:strLit><c:ptCount "
        f'val="{len(categories)}"/>{names}</c:strLit></c:cat>'
        f"<c:val><c:numLit><c:formatCode>General</c:formatCode>"
        f'<c:ptCount val="{len(stamps)}"/>{numbers}</c:numLit></c:val>'
        f"</c:ser>"
    )


def _series_style(index: int, kind: str) -> str:
    colors = ("35618D", "25847D", "B77936", "805EA0")
    color = colors[index % len(colors)]
    shape = (
        f'<c:spPr><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
        f'<a:ln w="19050"><a:solidFill><a:srgbClr val="{color}"/>'
        "</a:solidFill></a:ln></c:spPr>"
    )
    return shape + (
        '<c:marker><c:symbol val="none"/></c:marker>' if kind != "bar" else ""
    )


def chart_xml(node: NodeValue, timezone: str) -> bytes:
    """构造图表 XML，缺失值保留断点。Args: node, timezone。"""
    moments = sorted(
        {point.ts for series in node.series for point in series.points}
    )
    if not moments:
        raise ValueError("图表没有有效数据")
    categories = [
        moment.astimezone(ZoneInfo(timezone)).strftime("%m-%d\n%H:%M")
        for moment in moments
    ]
    stamps = [moment.isoformat() for moment in moments]
    series_xml = "".join(
        _series_xml(series, index, categories, stamps, node.kind)
        for index, series in enumerate(node.series)
    )
    kind = "barChart" if node.kind == "bar" else "lineChart"
    direction = '<c:barDir val="col"/>' if node.kind == "bar" else ""
    grouping = "clustered" if node.kind == "bar" else "standard"
    xml = (
        f'<c:chartSpace xmlns:c="{C_NS}" xmlns:a="{A_NS}"><c:chart>'
        '<c:autoTitleDeleted val="1"/><c:plotArea><c:layout/>'
        f'<c:{kind}>{direction}<c:grouping val="{grouping}"/>{series_xml}'
        f'<c:axId val="{_CAT_AXIS}"/><c:axId val="{_VAL_AXIS}"/></c:{kind}>'
        f'{_axis("catAx", _CAT_AXIS, _VAL_AXIS, "b", len(moments))}'
        f'{_axis("valAx", _VAL_AXIS, _CAT_AXIS, "l", len(moments))}'
        "</c:plotArea>"
        + (
            '<c:legend><c:legendPos val="b"/><c:layout/></c:legend>'
            if len(node.series) > 1
            else ""
        )
        + '<c:plotVisOnly val="1"/><c:dispBlanksAs val="gap"/>'
        '</c:chart><c:spPr><a:solidFill><a:srgbClr val="FFFFFF"/>'
        "</a:solidFill><a:ln><a:noFill/></a:ln></c:spPr>"
        '<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="900">'
        '<a:latin typeface="Noto Sans CJK SC"/>'
        '<a:ea typeface="Noto Sans CJK SC"/>'
        '</a:defRPr></a:pPr><a:endParaRPr lang="zh-CN"/>'
        "</a:p></c:txPr></c:chartSpace>"
    )
    parse_xml(xml.encode())
    return xml.encode()


def _axis(
    kind: str, axis_id: int, cross_id: int, position: str, count: int
) -> str:
    step = max(1, ceil(count / 4))
    category = (
        f'<c:tickLblSkip val="{step}"/><c:tickMarkSkip val="{step}"/>'
        if kind == "catAx"
        else ""
    )
    return (
        f'<c:{kind}><c:axId val="{axis_id}"/><c:scaling>'
        '<c:orientation val="minMax"/></c:scaling>'
        f'<c:delete val="0"/><c:axPos val="{position}"/>'
        '<c:numFmt formatCode="0.##" sourceLinked="0"/>'
        '<c:majorTickMark val="none"/><c:minorTickMark val="none"/>'
        '<c:tickLblPos val="nextTo"/>'
        f'<c:crossAx val="{cross_id}"/><c:crosses val="autoZero"/>'
        f"{category}</c:{kind}>"
    )


def add_chart(document: Document, node: NodeValue, timezone: str) -> None:
    """把图表部件挂到文档并内联显示。Args: document, node, timezone。"""
    package = document.part.package
    if package is None:
        raise ValueError("文档缺少 OPC 包")
    part = Part(
        package.next_partname("/word/charts/chart%d.xml"),
        CHART_CONTENT_TYPE,
        chart_xml(node, timezone),
        package,
    )
    relation = document.part.relate_to(part, CHART_REL)
    drawing_id = len(document.inline_shapes) + 1
    width = (
        int(document.sections[0].page_width or 7_560_000)
        - int(document.sections[0].left_margin or 0)
        - int(document.sections[0].right_margin or 0)
    )
    xml = (
        "<w:drawing "
        'xmlns:w="http://schemas.openxmlformats.org/'
        'wordprocessingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/'
        'drawingml/2006/wordprocessingDrawing" '
        f'xmlns:a="{A_NS}" xmlns:c="{C_NS}" xmlns:r="{R_NS}">'
        f'<wp:inline><wp:extent cx="{width}" cy="2800000"/>'
        f'<wp:docPr id="{drawing_id}" name="Report chart {drawing_id}"/>'
        f'<a:graphic><a:graphicData uri="{C_NS}"><c:chart r:id="{relation}"/>'
        "</a:graphicData></a:graphic></wp:inline></w:drawing>"
    )
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True
    append_xml(paragraph.add_run(), xml)
