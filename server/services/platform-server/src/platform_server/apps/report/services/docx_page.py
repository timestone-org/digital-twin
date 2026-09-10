"""Word 纸张、页眉页脚、水印及目录设置。"""

from xml.sax.saxutils import escape

from docx.document import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm

from platform_server.apps.report.schemas.document import PageSettings
from platform_server.apps.report.services.ooxml import append_xml
from platform_server.apps.report.services.typography import apply_typography
from platform_server.apps.report.services.watermark import watermark_xml

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

_PAPERS = {
    "A4": (21.0, 29.7),
    "A3": (29.7, 42.0),
    "Letter": (21.59, 27.94),
    "Legal": (21.59, 35.56),
}


def apply_page(document: Document, page: PageSettings) -> None:
    """设置纸张和正文默认样式。Args: document, page。"""
    section = document.sections[0]
    width, height = _PAPERS.get(page.size, (page.width_cm, page.height_cm))
    if page.orientation == "landscape":
        width, height = max(width, height), min(width, height)
        section.orientation = WD_ORIENT.LANDSCAPE
    if (
        page.margins_cm.left + page.margins_cm.right >= width
        or page.margins_cm.top + page.margins_cm.bottom >= height
    ):
        raise ValueError("页边距超过纸张可用范围")
    section.page_width, section.page_height = Cm(width), Cm(height)
    for side in ("top", "bottom", "left", "right"):
        setattr(section, f"{side}_margin", Cm(getattr(page.margins_cm, side)))
    section.header.paragraphs[0].text = page.header
    section.footer.paragraphs[0].text = page.footer
    apply_typography(document, page)
    if page.watermark and page.watermark.text:
        _watermark(document, page)


def _watermark(document: Document, page: PageSettings) -> None:
    watermark = page.watermark
    if watermark is None:
        return
    append_xml(
        document.sections[0].header.paragraphs[0].add_run(),
        watermark_xml(watermark),
    )


def add_toc(document: Document, headings: list[str]) -> None:
    """写入有缓存的目录。Args: document, headings。"""
    if not headings:
        return
    document.add_paragraph("目录", style="Caption")
    instruction = escape('TOC \\o "1-3" \\h \\z \\u')
    first = document.add_paragraph()
    append_xml(
        first.add_run(),
        f'<w:fldChar xmlns:w="{_W_NS}" w:fldCharType="begin"/>',
    )
    append_xml(
        first.add_run(),
        f'<w:instrText xmlns:w="{_W_NS}" xml:space="preserve">'
        f"{instruction}</w:instrText>",
    )
    append_xml(
        first.add_run(),
        f'<w:fldChar xmlns:w="{_W_NS}" w:fldCharType="separate"/>',
    )
    paragraph = first
    for index, text in enumerate(headings):
        paragraph = first if index == 0 else document.add_paragraph()
        paragraph.add_run(text)
        paragraph.paragraph_format.space_after = Cm(0.05)
    append_xml(
        paragraph.add_run(),
        f'<w:fldChar xmlns:w="{_W_NS}" w:fldCharType="end"/>',
    )
    append_xml(
        document.settings,
        '<w:updateFields xmlns:w="http://schemas.openxmlformats.org/'
        'wordprocessingml/2006/main" w:val="true"/>',
    )
