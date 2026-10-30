"""有界 Word 导入，映射已支持的格式并列出内容与版式损失。"""

import io
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from docx.document import Document as WordDocument
from docx.enum.section import WD_ORIENT
from docx.styles.style import ParagraphStyle
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from platform_server.apps.report.schemas.document import (
    DocumentNode,
    Margins,
    PageSettings,
)
from platform_server.apps.report.schemas.jobs import ImportResultOut
from platform_server.apps.report.services.docx_import_defaults import (
    DocumentDefaults,
    document_defaults,
)
from platform_server.apps.report.services.docx_import_styles import (
    paragraph_attrs,
    run_marks,
)
from platform_server.apps.report.services.docx_tables import import_table

MAX_IMPORT_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_ENTRIES = 2000
MAX_RELATIONSHIP_BYTES = 1024 * 1024
EMU_PER_CM = 360000


class WordImportInvalid(ValueError):
    """导入器明确校验过且可向用户显示的文件问题。"""


def validate_archive(payload: bytes) -> None:
    """拒绝超限压缩包与伪装的宏、主动内容。Args: payload。"""
    if len(payload) > MAX_IMPORT_BYTES:
        raise WordImportInvalid("Word 文件超过 10 MB")
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        entries = archive.infolist()
        if (
            len(entries) > MAX_ENTRIES
            or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES
        ):
            raise WordImportInvalid("Word 解压后体积超过上限")
        names = archive.namelist()
        if "word/document.xml" not in names:
            raise WordImportInvalid("不是有效的 Word 文档")
        if any(
            name.casefold().endswith("vbaproject.bin")
            or "/activex/" in name.casefold()
            for name in names
        ):
            raise WordImportInvalid("不支持带宏或主动内容的 Word 文档")
        if (
            "[Content_Types].xml" in names
            and b"macroenabled" in archive.read("[Content_Types].xml").lower()
        ):
            raise WordImportInvalid("不支持带宏的 Word 文档")


def import_docx(payload: bytes) -> ImportResultOut:
    """转换正文、格式和页面，并如实报告转换边界。Args: payload。"""
    validate_archive(payload)
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        dropped = _package_warnings(archive)
    document = Document(io.BytesIO(payload))
    defaults = document_defaults(document)
    converter = _Converter(dropped, defaults)
    content = [
        converter.block(block) for block in document.iter_inner_content()
    ]
    if len(document.sections) > 1:
        dropped.append("多节页面设置仅保留第一节")
    if document.inline_shapes:
        dropped.append(f"图片或图表 {len(document.inline_shapes)} 个未导入")
    return ImportResultOut(
        doc_json=DocumentNode(type="doc", content=content),
        page_json=_page(document, defaults),
        dropped=list(dict.fromkeys(dropped)),
    )


def _page(document: WordDocument, defaults: DocumentDefaults) -> PageSettings:
    section = document.sections[0]
    margins = Margins(
        **{
            side: round(
                float(getattr(section, f"{side}_margin") or 0) / EMU_PER_CM, 4
            )
            for side in ("top", "bottom", "left", "right")
        }
    )
    normal = document.styles["Normal"]
    font = normal.font if isinstance(normal, ParagraphStyle) else None
    size = (font.size if font else None) or defaults.font.size
    return PageSettings(
        size="custom",
        width_cm=float(section.page_width or 7560000) / EMU_PER_CM,
        height_cm=float(section.page_height or 10692000) / EMU_PER_CM,
        orientation=(
            "landscape"
            if section.orientation == WD_ORIENT.LANDSCAPE
            else "portrait"
        ),
        margins_cm=margins,
        header="\n".join(item.text for item in section.header.paragraphs),
        footer="\n".join(item.text for item in section.footer.paragraphs),
        font_family=(font.name if font else None)
        or defaults.font.name
        or "宋体",
        font_size_pt=min(72, max(6, size.pt)) if size is not None else 12,
    )


class _Converter:
    """一份导入内共享的损失清单。"""

    def __init__(self, warnings: list[str], defaults: DocumentDefaults) -> None:
        self.warnings = warnings
        self.defaults = defaults

    def block(self, block: Paragraph | Table) -> DocumentNode:
        if isinstance(block, Paragraph):
            return self.paragraph(block)
        self.warnings.append(
            "表格列宽、边框、底纹、单元格留白与行高采用报告样式，请核对版式"
        )
        return import_table(block, self.block)

    def paragraph(self, paragraph: Paragraph) -> DocumentNode:
        content: list[DocumentNode] = []
        for inline in paragraph.iter_inner_content():
            if isinstance(inline, Hyperlink):
                self.warnings.append("超链接已转为普通文本，未复制外部关系")
                for run in inline.runs:
                    content.extend(self.run(run, paragraph))
            else:
                content.extend(self.run(inline, paragraph))
        attrs = paragraph_attrs(paragraph, self.warnings, self.defaults)
        style_name = (
            (paragraph.style.name or "") if paragraph.style is not None else ""
        )
        if style_name.startswith("List"):
            self.warnings.append("列表自动编号转为普通段落")
        heading = (
            style_name
            and style_name.startswith("Heading ")
            and style_name[-1].isdigit()
        )
        if heading:
            attrs["level"] = min(6, int(style_name[-1]))
        return DocumentNode(
            type="heading" if heading else "paragraph",
            attrs=attrs,
            content=content,
        )

    def run(self, run: Run, paragraph: Paragraph) -> list[DocumentNode]:
        marks = run_marks(run, paragraph, self.warnings, self.defaults)
        content: list[DocumentNode] = []
        for index, text in enumerate(run.text.split("\n")):
            if index:
                content.append(DocumentNode(type="hardBreak"))
            if text:
                content.append(
                    DocumentNode(type="text", text=text, marks=marks)
                )
        return content


def _package_warnings(archive: zipfile.ZipFile) -> list[str]:
    warnings: list[str] = []
    names = archive.namelist()
    if any(
        name.startswith(("word/footnotes", "word/endnotes", "word/comments"))
        for name in names
    ):
        warnings.append("脚注、尾注及批注未导入")
    if any(name.startswith(("word/header", "word/footer")) for name in names):
        warnings.append("页眉页脚仅保留文字，原有格式、域和图形需重新设置")
    xml = archive.read("word/document.xml")
    for marker, message in (
        (b"w:numPr", "列表自动编号转为普通段落"),
        (b"w:fldChar", "动态域已转为当前显示文本，自动更新未保留"),
        (b"w:fldSimple", "简单域的自动更新及复杂内容未导入，请核对显示文本"),
        (b"w:drawing", "图片、图表及绘图对象未导入"),
        (b"w:pict", "形状、文本框及旧式绘图未导入"),
        (b"m:oMath", "Word 公式未导入"),
        (b"w:sdt", "内容控件未导入，请核对控件内文字"),
        (b"w:bookmarkStart", "书签与交叉引用关系未保留"),
        (b"w:ins", "修订标记未导入，请核对修订内容"),
        (b"w:del", "修订标记未导入，请核对修订内容"),
        (b'w:type="page"', "显式分页符未导入，请重新核对分页"),
    ):
        if marker in xml:
            warnings.append(message)
    for name in names:
        if name.endswith(".rels"):
            warnings.extend(_relationship_warnings(archive.read(name)))
    return warnings


def _relationship_warnings(payload: bytes) -> list[str]:
    if len(payload) > MAX_RELATIONSHIP_BYTES:
        raise WordImportInvalid("Word 关系表超过大小上限")
    scan = payload.replace(b"\x00", b"").upper()
    if b"<!DOCTYPE" in scan or b"<!ENTITY" in scan:
        raise WordImportInvalid("Word 关系表包含不安全 XML 声明")
    root = ET.fromstring(payload)  # noqa: S314  # 有界且已拒绝 DTD 与 ENTITY
    warnings: list[str] = []
    for relation in root:
        if relation.attrib.get("TargetMode", "").casefold() != "external":
            continue
        kind = relation.attrib.get("Type", "")
        warnings.append(
            "超链接已转为普通文本，未复制外部关系"
            if kind.endswith("/hyperlink")
            else "外部图片、模板或其他关系未导入，未访问外部地址"
        )
    return warnings
