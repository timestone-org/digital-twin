"""DOCX 的 PDF 派生：隔离 Office 子进程、校验产物并尽力落私有对象。"""

import asyncio
import uuid
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import parse_xml

from knowledge_server.apps.knowledge.services import docx_preview
from knowledge_server.apps.knowledge.services.docx_preview import (
    LibreOfficeDocxPreviewer,
    PreviewArtifact,
    PreviewRenderFailed,
    build_previewer,
    generate_preview,
)
from knowledge_server.apps.knowledge.services.parsing import RawItem
from knowledge_server.apps.knowledge.services.sources import preview_key
from lib.testing import FakeObjectStore

BASE = uuid.UUID("00000000-0000-7000-8000-000000000001")
DOC = uuid.UUID("00000000-0000-7000-8000-000000000002")
PDF = b"%PDF-1.7\nshape-preview"


class _Process:
    """`asyncio.subprocess.Process` 的窄替身。"""

    def __init__(
        self, *, is_hanging: bool = False, returncode: int = 0
    ) -> None:
        self.returncode: int | None = None if is_hanging else returncode
        self.is_hanging = is_hanging
        self.was_killed = False
        self.started = asyncio.Event()

    async def communicate(self) -> tuple[bytes, bytes]:
        self.started.set()
        if self.is_hanging:
            await asyncio.Future[None]()
        return (b"", b"")

    def kill(self) -> None:
        self.was_killed = True
        self.returncode = -9

    async def wait(self) -> int:
        return self.returncode or 0


def _docx(*, has_macro: bool = False) -> bytes:
    """现造一个带 Word 图形与外链关系的有效 DOCX。

    Args: has_macro。
    """
    document = Document()
    document.add_paragraph("REAL-CONVERSION-SENTINEL")
    paragraph = document.add_paragraph()
    paragraph._p.append(
        parse_xml(
            """<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:rect id="shape1"
 style="width:240pt;height:120pt" fillcolor="#ff0000" strokecolor="#000000">
 <v:textbox><w:txbxContent><w:p><w:r><w:t>SHAPE-SENTINEL</w:t></w:r></w:p>
 </w:txbxContent></v:textbox></v:rect></w:pict></w:r>"""
        )
    )
    document.part.relate_to(
        "http://169.254.169.254/latest/meta-data",
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    target = BytesIO()
    document.save(target)
    if has_macro:
        with zipfile.ZipFile(target, "a") as made:
            made.writestr("word/vbaProject.bin", b"macro")
    return target.getvalue()


def _raw(name: str = "手册.docx", *, has_macro: bool = False) -> RawItem:
    return RawItem(
        filename=name, media_type="", content=_docx(has_macro=has_macro)
    )


async def test_a_shape_docx_becomes_a_sanitized_pdf(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """⚠ Office 子进程不能继承服务密钥或跟随 DOCX 外链。"""
    captured: dict[str, object] = {}

    async def launch(*command: str, **options: object) -> _Process:
        captured["command"] = command
        captured["env"] = options.get("env")
        source = Path(command[-1])
        output_dir = Path(command[command.index("--outdir") + 1])
        with zipfile.ZipFile(source) as opened:
            rels = opened.read("word/_rels/document.xml.rels")
            body = opened.read("word/document.xml")
        assert b'TargetMode="External"' not in rels
        assert b"xmlns:ns0" not in rels
        assert b"SHAPE-SENTINEL" in body
        (output_dir / "document.pdf").write_bytes(PDF)
        return _Process()

    monkeypatch.setattr(
        docx_preview.shutil, "which", lambda _one: "/bin/soffice"
    )
    monkeypatch.setattr(docx_preview.asyncio, "create_subprocess_exec", launch)
    monkeypatch.setenv("KNOWLEDGE_OBJECTSTORE_SECRET_KEY", "must-not-leak")

    made = await LibreOfficeDocxPreviewer(
        executable="soffice", timeout_s=1
    ).render(_raw())

    assert made == PreviewArtifact(content=PDF)
    command = captured["command"]
    assert isinstance(command, tuple)
    assert "--safe-mode" in command
    assert any(
        str(one).startswith("-env:UserInstallation=file:") for one in command
    )
    environment = captured["env"]
    assert isinstance(environment, dict)
    assert "KNOWLEDGE_OBJECTSTORE_SECRET_KEY" not in environment


async def test_non_docx_never_starts_libreoffice(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def forbidden(*_args: object, **_kwargs: object) -> _Process:
        raise AssertionError("不该启动子进程")

    monkeypatch.setattr(
        docx_preview.asyncio, "create_subprocess_exec", forbidden
    )
    made = await LibreOfficeDocxPreviewer(
        executable="soffice", timeout_s=1
    ).render(_raw("手册.md"))
    assert made is None


async def test_a_macro_disguised_as_docx_is_not_opened(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def forbidden(*_args: object, **_kwargs: object) -> _Process:
        raise AssertionError("不该启动子进程")

    monkeypatch.setattr(
        docx_preview.asyncio, "create_subprocess_exec", forbidden
    )
    with pytest.raises(PreviewRenderFailed, match="宏"):
        await LibreOfficeDocxPreviewer(
            executable="soffice", timeout_s=1
        ).render(_raw(has_macro=True))


@pytest.mark.parametrize(
    "output",
    [b"not a pdf", PDF + b"x" * 100],
    ids=["wrong-type", "too-large"],
)
async def test_invalid_converter_output_is_rejected(
    monkeypatch: pytest.MonkeyPatch, output: bytes
) -> None:
    async def launch(*command: str, **_options: object) -> _Process:
        output_dir = Path(command[command.index("--outdir") + 1])
        (output_dir / "document.pdf").write_bytes(output)
        return _Process()

    monkeypatch.setattr(
        docx_preview.shutil, "which", lambda _one: "/bin/soffice"
    )
    monkeypatch.setattr(docx_preview.asyncio, "create_subprocess_exec", launch)
    with pytest.raises(PreviewRenderFailed):
        await LibreOfficeDocxPreviewer(
            executable="soffice", timeout_s=1, max_output_bytes=len(PDF)
        ).render(_raw())


@pytest.mark.parametrize(
    "content",
    [
        b"<!DOCTYPE Relationships><Relationships/>",
        "<!DOCTYPE Relationships><Relationships/>".encode("utf-16"),
        "<!DOCTYPE Relationships><Relationships/>".encode("utf-16-le"),
        "<!DOCTYPE Relationships><Relationships/>".encode("utf-16-be"),
    ],
    ids=["utf-8", "utf-16-bom", "utf-16-le", "utf-16-be"],
)
def test_relationship_dtd_is_rejected_in_every_supported_encoding(
    content: bytes,
) -> None:
    with pytest.raises(PreviewRenderFailed, match="不安全"):
        docx_preview._safe_relationships(content)


def test_an_oversized_relationship_part_is_rejected_before_parsing() -> None:
    content = b" " * (docx_preview.MAX_RELATIONSHIP_BYTES + 1)
    with pytest.raises(PreviewRenderFailed, match="过大"):
        docx_preview._safe_relationships(content)


@pytest.mark.parametrize(
    "content",
    [b"\xff", b"<Relationships>"],
    ids=["bad-encoding", "bad-xml"],
)
def test_malformed_relationship_xml_is_rejected(content: bytes) -> None:
    with pytest.raises(PreviewRenderFailed, match="关系表"):
        docx_preview._safe_relationships(content)


async def test_a_file_that_is_not_a_zip_is_rejected() -> None:
    with pytest.raises(PreviewRenderFailed, match="压缩结构"):
        await LibreOfficeDocxPreviewer(
            executable="soffice", timeout_s=1
        ).render(RawItem(filename="fake.docx", media_type="", content=b"no"))


async def test_a_success_exit_without_a_pdf_is_rejected(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def launch(*_args: object, **_options: object) -> _Process:
        return _Process()

    monkeypatch.setattr(
        docx_preview.shutil, "which", lambda _one: "/bin/soffice"
    )
    monkeypatch.setattr(docx_preview.asyncio, "create_subprocess_exec", launch)
    with pytest.raises(PreviewRenderFailed, match="没有产出"):
        await LibreOfficeDocxPreviewer(
            executable="soffice", timeout_s=1
        ).render(_raw())


async def test_a_nonzero_libreoffice_exit_is_rejected(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def launch(*_args: object, **_options: object) -> _Process:
        return _Process(returncode=1)

    monkeypatch.setattr(
        docx_preview.shutil, "which", lambda _one: "/bin/soffice"
    )
    monkeypatch.setattr(docx_preview.asyncio, "create_subprocess_exec", launch)
    with pytest.raises(PreviewRenderFailed, match="无法转换"):
        await LibreOfficeDocxPreviewer(
            executable="soffice", timeout_s=1
        ).render(_raw())


async def test_a_hung_converter_is_killed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    process = _Process(is_hanging=True)

    async def launch(*_args: object, **_options: object) -> _Process:
        return process

    monkeypatch.setattr(
        docx_preview.shutil, "which", lambda _one: "/bin/soffice"
    )
    monkeypatch.setattr(docx_preview.asyncio, "create_subprocess_exec", launch)
    with pytest.raises(PreviewRenderFailed, match="超时"):
        await LibreOfficeDocxPreviewer(
            executable="soffice", timeout_s=0.01
        ).render(_raw())
    assert process.was_killed is True


async def test_cancelling_conversion_kills_the_child_process(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    process = _Process(is_hanging=True)

    async def launch(*_args: object, **_options: object) -> _Process:
        return process

    monkeypatch.setattr(
        docx_preview.shutil, "which", lambda _one: "/bin/soffice"
    )
    monkeypatch.setattr(docx_preview.asyncio, "create_subprocess_exec", launch)
    task = asyncio.create_task(
        LibreOfficeDocxPreviewer(executable="soffice", timeout_s=30).render(
            _raw()
        )
    )
    await process.started.wait()

    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task

    assert process.was_killed is True


@dataclass(frozen=True)
class _Previewer:
    """给存储编排用的假派生器。"""

    is_failing: bool = False
    is_empty: bool = False

    async def probe(self) -> None:
        return None

    async def render(self, raw: RawItem) -> PreviewArtifact | None:
        del raw
        if self.is_failing:
            raise PreviewRenderFailed("坏了")
        if self.is_empty:
            return None
        return PreviewArtifact(PDF)


async def test_generated_pdf_uses_the_document_private_prefix() -> None:
    store = FakeObjectStore()
    assert await generate_preview(_Previewer(), store, (BASE, DOC), _raw())
    assert store.objects[preview_key(BASE, DOC)] == (PDF, "application/pdf")


async def test_preview_failure_does_not_fail_document_ingest() -> None:
    store = FakeObjectStore()
    made = await generate_preview(
        _Previewer(is_failing=True), store, (BASE, DOC), _raw()
    )
    assert made is False
    assert store.objects == {}


async def test_absent_previewer_or_irrelevant_format_writes_nothing() -> None:
    store = FakeObjectStore()
    assert not await generate_preview(None, store, (BASE, DOC), _raw())
    assert not await generate_preview(
        _Previewer(is_empty=True), store, (BASE, DOC), _raw("手册.md")
    )
    assert store.objects == {}


async def test_enabled_previewer_requires_an_installed_command(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(docx_preview.shutil, "which", lambda _one: None)
    with pytest.raises(PreviewRenderFailed, match="找不到 LibreOffice"):
        await LibreOfficeDocxPreviewer(
            executable="missing", timeout_s=1
        ).probe()


def test_the_feature_switch_controls_previewer_assembly() -> None:
    assert build_previewer(False, "soffice", 30) is None
    assert isinstance(
        build_previewer(True, "soffice", 30), LibreOfficeDocxPreviewer
    )
